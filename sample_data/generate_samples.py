from __future__ import annotations

import json
from pathlib import Path

import pymupdf
from docx import Document


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "generated"

SAMPLES = {
    "ananya_shah_python_developer": """Ananya Shah
ananya.shah@example.com | +91 90000 00001

SUMMARY
Python developer with 3 years of experience building internal web tools and REST APIs.

TECHNICAL SKILLS
Python, FastAPI, Flask, SQL, PostgreSQL, Git, Docker, Azure App Service

EXPERIENCE
Software Developer — Northstar Labs (2023–Present)
Built tested REST APIs in Python and PostgreSQL and containerized services with Docker.

EDUCATION
B.Tech in Computer Science — Fictional Institute of Technology, 2023

PROJECTS
Created a Streamlit issue-triage dashboard with Plotly.

CERTIFICATIONS
Microsoft Certified: Azure Fundamentals
""",
    "kabir_mehta_data_analyst": """Kabir Mehta
kabir.mehta@example.com | +91 90000 00002

SUMMARY
Data analyst with 2 years of experience producing operational dashboards.

TECHNICAL SKILLS
SQL, Microsoft Excel, Power BI, Tableau, Python, pandas, statistics, data visualization

EXPERIENCE
Data Analyst — Bluebird Retail (2024–Present)
Prepared SQL datasets and Power BI dashboards; documented quality checks for stakeholders.

EDUCATION
B.Sc. in Statistics — Example University, 2024

PROJECTS
Analyzed fictional retail trends using pandas and Plotly.

CERTIFICATIONS
Microsoft Certified: Power BI Data Analyst Associate
""",
    "leena_iyer_cloud_engineer": """Leena Iyer
leena.iyer@example.com | +91 90000 00003

SUMMARY
Cloud engineer with 4 years of experience supporting secure cloud platforms.

TECHNICAL SKILLS
Microsoft Azure, Azure Functions, Azure Blob Storage, Docker, Kubernetes, Terraform, Linux, CI/CD, Git

EXPERIENCE
Cloud Engineer — Contoso Learning Lab (2022–Present)
Provisioned private storage and App Service environments using Terraform and CI/CD pipelines.

EDUCATION
B.E. in Information Technology — Sample College, 2022

PROJECTS
Designed a private Blob Storage upload workflow with retention controls.

CERTIFICATIONS
Microsoft Certified: Azure Administrator Associate
""",
    "omar_khan_ml_engineer": """Omar Khan
omar.khan@example.com | +91 90000 00004

SUMMARY
AI/ML engineer with 3.5 years of experience developing human-reviewed NLP prototypes.

TECHNICAL SKILLS
Python, PyTorch, scikit-learn, natural language processing, machine learning, pandas, REST API, Docker, Azure

EXPERIENCE
Machine Learning Engineer — Responsible AI Studio (2023–Present)
Built NLP evaluation pipelines and evidence-linked model reports with scikit-learn and PyTorch.

EDUCATION
M.Sc. in Data Science — Demonstration University, 2023

PROJECTS
Created an explainable resume parsing prototype using NLP and deterministic scoring.

CERTIFICATIONS
Microsoft Certified: Azure AI Engineer Associate
""",
}

JOBS = [
    {
        "title": "Python Developer",
        "description": "Build Python REST APIs, work with SQL and Git, and deploy tested services using Docker.",
        "required_skills": ["Python", "REST API", "SQL", "Git"],
        "preferred_skills": ["FastAPI", "Docker", "Azure App Service"],
        "minimum_experience_years": 2,
    },
    {
        "title": "Data Analyst",
        "description": "Create trustworthy business reports using SQL, Excel, Power BI, statistics, and data visualization.",
        "required_skills": ["SQL", "Excel", "Power BI", "data visualization"],
        "preferred_skills": ["Python", "pandas", "Tableau"],
        "minimum_experience_years": 1,
    },
    {
        "title": "Cloud Engineer",
        "description": "Operate Azure services using Terraform, Linux, containers, CI/CD, and private storage practices.",
        "required_skills": ["Azure", "Terraform", "Linux", "Docker", "CI/CD"],
        "preferred_skills": ["Kubernetes", "Azure Functions", "Azure Blob Storage"],
        "minimum_experience_years": 3,
    },
    {
        "title": "AI/ML Engineer",
        "description": "Develop Python machine-learning and NLP systems with scikit-learn or PyTorch and clear evaluation evidence.",
        "required_skills": ["Python", "machine learning", "natural language processing"],
        "preferred_skills": ["scikit-learn", "PyTorch", "Docker", "Azure"],
        "minimum_experience_years": 2,
    },
]


def make_docx(path: Path, text: str) -> None:
    document = Document()
    for index, line in enumerate(text.splitlines()):
        if index == 0:
            document.add_heading(line, level=0)
        elif line.isupper() and line:
            document.add_heading(line.title(), level=1)
        elif line:
            document.add_paragraph(line)
    document.save(path)


def make_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    y = 55
    for index, line in enumerate(text.splitlines()):
        if not line:
            y += 8
            continue
        size = 16 if index == 0 else 11
        if line.isupper():
            size = 12
        if y > 790:
            page = document.new_page(width=595, height=842)
            y = 55
        page.insert_text((55, y), line, fontsize=size, fontname="helv")
        y += size + 7
    document.save(path)
    document.close()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for name, text in SAMPLES.items():
        make_docx(OUTPUT / f"{name}.docx", text)
        make_pdf(OUTPUT / f"{name}.pdf", text)
    (ROOT / "sample_jobs.json").write_text(json.dumps(JOBS, indent=2), encoding="utf-8")
    print(f"Created {len(SAMPLES) * 2} fictional resumes in {OUTPUT}")


if __name__ == "__main__":
    main()

