from __future__ import annotations

from pathlib import Path


class DocumentParseError(ValueError):
    pass


class DocumentParser:
    def extract(self, path: str | Path) -> str:
        path = Path(path)
        try:
            if path.suffix.lower() == ".pdf":
                return self._extract_pdf(path)
            if path.suffix.lower() == ".docx":
                return self._extract_docx(path)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError(f"Could not parse {path.name}: {exc}") from exc
        raise DocumentParseError("Only PDF and DOCX files are supported.")

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            import pymupdf
        except ImportError as exc:
            raise DocumentParseError("PyMuPDF is not installed.") from exc
        try:
            with pymupdf.open(path) as document:
                if document.page_count == 0:
                    raise DocumentParseError("The PDF has no pages.")
                text = "\n".join(page.get_text("text") for page in document)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("The PDF is corrupt, encrypted, or unreadable.") from exc
        if not text.strip():
            raise DocumentParseError("No selectable text was found. Enable Azure OCR for a scanned resume.")
        return text.strip()

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentParseError("python-docx is not installed.") from exc
        try:
            document = Document(path)
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(part for part in parts if part.strip())
        except Exception as exc:
            raise DocumentParseError("The DOCX file is corrupt or unreadable.") from exc
        if not text.strip():
            raise DocumentParseError("The DOCX file contains no readable text.")
        return text.strip()
