from __future__ import annotations

import json
import sqlite3
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from components.candidate_view import candidate_report
from config.settings import AppConfig
from database.repository import Repository
from models.domain import (
    CandidateProfile,
    EmploymentEntry,
    EvidenceSnippet,
    JobDescription,
    JobRequirement,
    ProcessingStatus,
    QualityStatus,
    RequirementCategory,
    RequirementEvaluationMethod,
    RequirementImportance,
    RequirementStatus,
)
from services.auth_service import AccessContext
from services.correction_service import CandidateCorrectionService
from services.export_service import export_csv, export_excel, results_frame
from services.matching_service import MatchingService
from services.privacy_service import PrivacyService
from services.quality_service import JobQualityService
from services.requirement_service import RequirementService
from services.skill_service import SemanticSimilarityService, SkillService


def _matcher() -> MatchingService:
    return MatchingService(SkillService(), SemanticSimilarityService("tfidf"))


def _candidate() -> CandidateProfile:
    return CandidateProfile(
        id=7,
        file_name="Jordan-Kapoor-Resume.docx",
        file_hash="secret-document-hash",
        candidate_name="Jordan Kapoor",
        email="jordan@example.test",
        phone="+91 98765 43210",
        skills=["python"],
        skill_evidence={
            "python": [
                EvidenceSnippet(
                    qualification="python",
                    snippet="Jordan Kapoor used Python at Secret Labs, 12 Park Road; portfolio https://example.test/jordan.",
                )
            ]
        },
        employment_history=[
            EmploymentEntry(
                title="Developer",
                organization="Secret Labs",
                period="2022-2024",
                evidence="Jordan Kapoor was a Developer at Secret Labs.",
            )
        ],
        raw_text="Jordan Kapoor jordan@example.test +91 98765 43210",
        processing_status=ProcessingStatus.PROCESSED,
    )


def _job() -> JobDescription:
    return JobDescription(
        id=3,
        title="Python Developer",
        description="Build reliable Python services, test integrations, review evidence, and collaborate with the engineering team.",
        required_skills=["python"],
    )


def test_anonymous_identity_is_stable_and_exports_are_sanitized() -> None:
    privacy = PrivacyService()
    candidate = _candidate()
    job = _job()
    result = _matcher().analyze(candidate, job)

    anonymous = privacy.anonymize_candidate(candidate)
    anonymous_again = privacy.anonymize_candidate(candidate)
    anonymous_result = privacy.anonymize_match(result, candidate)
    assert anonymous.candidate_name == anonymous_again.candidate_name == "Candidate 007"
    assert anonymous.id == candidate.id
    assert anonymous.email is None and anonymous.phone is None
    assert anonymous.file_name == "candidate-007-resume.docx"
    assert anonymous.file_hash == "anonymous-7"
    assert "Jordan" not in anonymous_result.model_dump_json()
    assert "Secret Labs" not in anonymous_result.model_dump_json()
    assert "12 Park Road" not in anonymous_result.model_dump_json()
    assert candidate.candidate_name == "Jordan Kapoor"  # the stored source model is unchanged

    frame = results_frame([anonymous_result], {candidate.id: anonymous}, {job.id: job})
    csv_bytes = export_csv(frame)
    json_bytes = candidate_report(anonymous, anonymous_result)
    excel_bytes = export_excel(frame, pd.DataFrame([{"event_type": "review"}]))
    workbook = load_workbook(BytesIO(excel_bytes), read_only=True)
    excel_text = " ".join(
        str(cell.value or "")
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    for forbidden in ["Jordan Kapoor", "jordan@example.test", "+91 98765 43210", "Secret Labs", "secret-document-hash"]:
        assert forbidden.encode() not in csv_bytes
        assert forbidden.encode() not in json_bytes
        assert forbidden not in excel_text


def test_reveal_authorization_and_audit(tmp_path: Path) -> None:
    repository = Repository(tmp_path / "reveal.sqlite3")
    candidate = repository.add_candidate(_candidate().model_copy(update={"id": None}))
    job = repository.add_job(_job().model_copy(update={"id": None}))
    reviewer = AccessContext(True, "Reviewer", user_id="reviewer-1", roles=("reviewer",))
    viewer = AccessContext(True, "Viewer", user_id="viewer-1", roles=("viewer",))

    PrivacyService.reveal_identity(repository, candidate.id, job.id, reviewer, False)
    event = repository.list_audit_events(limit=1)[0]
    assert event["event_type"] == "candidate_identity_revealed"
    assert event["actor"] == "reviewer-1"
    assert event["details"] == {"job_id": job.id}
    with pytest.raises(PermissionError):
        PrivacyService.reveal_identity(repository, candidate.id, job.id, viewer, False)
    with pytest.raises(PermissionError):
        PrivacyService.reveal_identity(repository, candidate.id, job.id, reviewer, True)


def test_demo_mode_forces_anonymous_review() -> None:
    assert AppConfig(demo_mode=True, anonymous_default=False).anonymous_review_enabled(False) is True
    assert AppConfig(demo_mode=False, anonymous_default=True).anonymous_review_enabled() is True


def test_correction_creation_recalculation_history_and_reversion(tmp_path: Path) -> None:
    repository = Repository(tmp_path / "corrections.sqlite3")
    candidate = repository.add_candidate(
        CandidateProfile(
            file_name="fictional.docx",
            file_hash="fictional-hash",
            candidate_name="Fictional Candidate",
            skills=["python"],
            raw_text="Python developer",
            processing_status=ProcessingStatus.PROCESSED,
        )
    )
    job = repository.add_job(
        JobDescription(title="Data Role", description="SQL data pipelines and reliable services", required_skills=["sql"])
    )
    matcher = _matcher()
    repository.save_match(matcher.analyze(candidate, job))
    service = CandidateCorrectionService(repository, matcher)
    access = AccessContext(True, "Reviewer", user_id="reviewer-9", roles=("reviewer",))

    correction, recalculated = service.correct(
        candidate.id,
        "skills",
        ["python", "sql"],
        "SQL was listed in the skills section.",
        access,
        False,
    )
    assert correction.original_value == ["python"]
    assert correction.corrected_value == ["python", "sql"]
    assert correction.actor == "reviewer-9"
    assert recalculated == 1
    assert repository.get_candidate(candidate.id).skills == ["python", "sql"]
    assert repository.get_match(candidate.id, job.id).required_coverage == 100
    assert len(repository.list_match_history(candidate.id, job.id)) == 2
    assert repository.list_candidate_corrections(candidate.id)[0].reason.startswith("SQL")
    assert repository.list_audit_events(limit=1)[0]["event_type"] == "candidate_correction_created"

    recalculated = service.revert(correction.id, access, False)
    assert recalculated == 1
    assert repository.get_candidate(candidate.id).skills == ["python"]
    assert repository.get_match(candidate.id, job.id).required_coverage == 0
    assert repository.get_correction(correction.id).active is False
    assert repository.list_audit_events(limit=1)[0]["event_type"] == "candidate_correction_reverted"


def test_correction_reason_and_service_authorization_are_enforced(tmp_path: Path) -> None:
    repository = Repository(tmp_path / "correction-auth.sqlite3")
    candidate = repository.add_candidate(_candidate().model_copy(update={"id": None}))
    service = CandidateCorrectionService(repository, _matcher())
    reviewer = AccessContext(True, "Reviewer", roles=("reviewer",))
    viewer = AccessContext(True, "Viewer", roles=("viewer",))
    with pytest.raises(ValueError, match="meaningful"):
        service.correct(candidate.id, "phone", "0000000000", "", reviewer, False)
    with pytest.raises(PermissionError):
        service.correct(candidate.id, "phone", "0000000000", "Verified", viewer, False)
    with pytest.raises(PermissionError):
        service.correct(candidate.id, "phone", "0000000000", "Verified", reviewer, True)


def test_quality_checker_classifies_blockers_warnings_suggestions_and_contradictions() -> None:
    checker = JobQualityService()
    required = JobRequirement(text="Python", normalized_value="python")
    duplicate = required.model_copy()
    excluded = required.model_copy(update={"importance": RequirementImportance.EXCLUDED})
    description = (
        "We need a young and energetic hard worker with Python and SQL skills who can build services, "
        "document decisions, collaborate with engineers, and deliver stable production changes."
    )
    result = checker.check(
        "Developer",
        description,
        [required, duplicate, excluded],
        minimum_experience_years=20,
    )
    codes = {item.code: item.severity.value for item in result.findings}
    assert result.status == QualityStatus.BLOCKED
    assert codes["vague_language"] in {"blocking", "warning"}
    assert codes["duplicate_requirement"] == "warning"
    assert codes["contradictory_requirement"] == "blocking"
    assert codes["excessive_experience"] == "warning"
    assert "unstructured_skill" in codes

    sensitive = checker.check("Developer", "This role requires Python and specifies religion for applicant screening.", [required])
    assert sensitive.status == QualityStatus.BLOCKED
    assert any(item.code == "sensitive_criterion" for item in sensitive.findings)
    empty = checker.check("", "Too short", [])
    assert empty.status == QualityStatus.BLOCKED
    assert {item.code for item in empty.findings} >= {"missing_title", "short_description", "no_required_requirements"}


def test_quality_warning_acknowledgement_is_required_and_audited(tmp_path: Path) -> None:
    repository = Repository(tmp_path / "quality.sqlite3")
    job = repository.add_job(_job().model_copy(update={"id": None}))
    checker = JobQualityService()
    requirement = JobRequirement(text="Python", normalized_value="python")
    result = checker.check(
        job.title,
        "Python is required. Good communication is expected while building tested and reliable services for users.",
        [requirement],
    )
    assert result.status == QualityStatus.REVIEW
    access = AccessContext(True, "Reviewer", user_id="reviewer-2", roles=("reviewer",))
    with pytest.raises(ValueError, match="acknowledgement"):
        checker.save_review(repository, result.model_copy(deep=True), job.id, False, access)
    saved = checker.save_review(repository, result, job.id, True, access)
    assert saved.warnings_acknowledged is True
    event = repository.list_audit_events(limit=1)[0]
    assert event["event_type"] == "quality_warnings_acknowledged"
    assert event["actor"] == "reviewer-2"


def test_requirement_versions_preserve_history_and_drive_matching(tmp_path: Path) -> None:
    repository = Repository(tmp_path / "requirements.sqlite3")
    service = RequirementService()
    job = repository.add_job(
        JobDescription(title="Developer", description="Python delivery, SQL context, and manual client communication review.")
    )
    first_rows = [
        JobRequirement(text="Python", normalized_value="python", sort_order=1),
        JobRequirement(
            text="SQL",
            normalized_value="sql",
            importance=RequirementImportance.EXCLUDED,
            enabled=True,
            sort_order=2,
        ),
        JobRequirement(
            text="Client communication",
            normalized_value="client communication",
            category=RequirementCategory.RESPONSIBILITY,
            importance=RequirementImportance.MANUAL,
            evaluation_method=RequirementEvaluationMethod.HUMAN_ONLY,
            sort_order=0,
        ),
        JobRequirement(text="Python", normalized_value="python", importance=RequirementImportance.PREFERRED, sort_order=3),
    ]
    merged, duplicate_count = service.merge_duplicates(first_rows)
    assert duplicate_count == 1
    assert [item.sort_order for item in merged] == [0, 1, 2]
    assert service.validate(merged) == []
    job, version_one, created = service.save_version(repository, job, merged, "reviewer")
    assert created is True and version_one.version_number == 1
    assert job.requirement_version_id == version_one.id

    second_rows = [*merged, JobRequirement(text="Docker", normalized_value="docker", importance=RequirementImportance.PREFERRED)]
    job, version_two, created = service.save_version(repository, job, second_rows, "reviewer", "Added Docker")
    assert created is True and version_two.version_number == 2
    assert [item.version_number for item in repository.list_requirement_versions(job.id)] == [2, 1]
    _, same_version, created = service.save_version(repository, job, second_rows, "reviewer")
    assert created is False and same_version.id == version_two.id

    candidate = repository.add_candidate(
        CandidateProfile(
            file_name="fictional.docx",
            file_hash="requirement-candidate",
            candidate_name="Fictional Candidate",
            skills=["python", "sql"],
            raw_text="Python SQL and client communication",
        )
    )
    result = _matcher().analyze(candidate, job)
    assert result.requirement_version_id == version_two.id
    assert result.required_coverage == 100
    assert not any(item.requirement.casefold() == "sql" for item in result.assessments)
    manual = next(item for item in result.assessments if item.requirement == "Client communication")
    assert manual.status == RequirementStatus.MANUAL_VERIFY

    protected = [JobRequirement(text="Religion", normalized_value="religion")]
    assert any("Protected" in error for error in service.validate(protected))


def test_additive_schema_migration_preserves_legacy_candidate(tmp_path: Path) -> None:
    database_path = tmp_path / "legacy.sqlite3"
    candidate = CandidateProfile(file_name="legacy.docx", file_hash="legacy", candidate_name="Legacy Candidate")
    with sqlite3.connect(database_path) as connection:
        connection.executescript(
            """
            CREATE TABLE candidates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                file_name TEXT NOT NULL,
                file_hash TEXT NOT NULL UNIQUE,
                status TEXT NOT NULL,
                candidate_name TEXT NOT NULL,
                payload TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            """
        )
        connection.execute(
            "INSERT INTO candidates(file_name,file_hash,status,candidate_name,payload,created_at) VALUES(?,?,?,?,?,?)",
            (
                candidate.file_name,
                candidate.file_hash,
                candidate.processing_status.value,
                candidate.candidate_name,
                candidate.model_dump_json(exclude={"reviewer_comments"}),
                candidate.created_at.isoformat(),
            ),
        )
    repository = Repository(database_path)
    assert repository.list_candidates()[0].candidate_name == "Legacy Candidate"
    assert repository.list_candidates()[0].reviewer_comments == []
    with repository.connect() as connection:
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"candidate_corrections", "job_requirement_versions", "quality_check_results", "match_history"} <= tables
